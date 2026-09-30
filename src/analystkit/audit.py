"""Conservative workbook checks with exact cell references."""

from dataclasses import dataclass, asdict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator, TranslatorError


@dataclass(frozen=True)
class AuditFinding:
    sheet: str
    cell: str
    severity: str
    category: str
    explanation: str


@dataclass
class AuditResult:
    source: str
    findings: list[AuditFinding]
    sheets_checked: int

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Audit exports to .xlsx")
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame([asdict(f) for f in self.findings], columns=list(AuditFinding.__dataclass_fields__)).to_excel(writer, sheet_name="Findings", index=False)
            pd.DataFrame([{"source": self.source, "sheets_checked": self.sheets_checked,
                           "findings": len(self.findings)}]).to_excel(writer, sheet_name="Summary", index=False)
        return path


def audit(path: str | Path) -> AuditResult:
    """Flag errors, broken references, hidden sheets, and formula outliers for review."""
    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Spreadsheet audit requires .xlsx or .xlsm")
    book = load_workbook(path, read_only=False, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    findings = []
    try:
        for sheet in book:
            if sheet.sheet_state != "visible":
                findings.append(AuditFinding(sheet.title, "", "review", "hidden sheet", "This sheet is hidden; check whether it affects the report."))
            formula_columns: dict[int, list] = {}
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "e":
                        findings.append(AuditFinding(sheet.title, cell.coordinate, "error", "cell error", f"Excel error: {cell.value}"))
                    if cell.data_type == "f":
                        formula_columns.setdefault(cell.column, []).append(cell)
                        if "#REF!" in str(cell.value).upper():
                            findings.append(AuditFinding(sheet.title, cell.coordinate, "error", "broken reference", "Formula contains #REF!."))
            for column, formulas in formula_columns.items():
                if len(formulas) < 3:
                    continue
                patterns: dict[str, list] = {}
                for cell in formulas:
                    try:
                        pattern = Translator(cell.value, origin=cell.coordinate).translate_formula(f"{cell.column_letter}100000")
                    except (TranslatorError, ValueError):
                        continue
                    patterns.setdefault(pattern, []).append(cell)
                dominant = max(patterns.values(), key=len, default=[])
                if len(dominant) >= 3 and len(dominant) > len(formulas) / 2:
                    for pattern, cells in patterns.items():
                        if cells is dominant:
                            continue
                        for cell in cells:
                            findings.append(AuditFinding(sheet.title, cell.coordinate, "review", "formula pattern", "Formula differs from the usual formula in this column; check if intentional."))
                first, last = min(c.row for c in formulas), max(c.row for c in formulas)
                if last - first > 5000:
                    continue
                for row_number in range(first, last + 1):
                    cell = sheet.cell(row_number, column)
                    if cell.data_type == "f" or cell.value is None:
                        continue
                    if sheet.cell(row_number, max(1, column - 1)).value is not None:
                        findings.append(AuditFinding(sheet.title, cell.coordinate, "review", "fixed value in formula column", "A fixed value appears inside a column that otherwise uses formulas."))
    finally:
        book.close()
    return AuditResult(path.name, findings, len(book.sheetnames))
