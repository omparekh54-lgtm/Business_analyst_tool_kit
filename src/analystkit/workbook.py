"""Find tables and headers without relying on sheet names or fixed row numbers."""

from dataclasses import dataclass, field
from pathlib import Path
import csv
import re

import pandas as pd
from openpyxl import load_workbook


@dataclass(frozen=True)
class TableCandidate:
    sheet: str
    header_row: int  # Excel's one-based row number
    columns: tuple[str, ...]
    data_rows: int
    score: float
    warnings: tuple[str, ...] = ()
    start_column: int = 1  # Excel's one-based column number


@dataclass
class WorkbookInspection:
    path: Path
    candidates: list[TableCandidate] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _header_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _unique_headers(values: tuple[object, ...]) -> tuple[tuple[str, ...], bool]:
    seen: dict[str, int] = {}
    output = []
    duplicates = False
    for index, value in enumerate(values, 1):
        name = _header_name(value) or f"Unnamed {index}"
        key = name.casefold()
        seen[key] = seen.get(key, 0) + 1
        if seen[key] > 1:
            duplicates = True
            name = f"{name} ({seen[key]})"
        output.append(name)
    return tuple(output), duplicates


def _candidate(rows: list[tuple[object, ...]], sheet: str, row_index: int, max_row: int) -> TableCandidate | None:
    row = rows[row_index]
    used = [i for i, value in enumerate(row) if value is not None and str(value).strip()]
    if len(used) < 2:
        return None
    start, end = min(used), max(used) + 1
    cells = row[start:end]
    text_count = sum(isinstance(v, str) and bool(v.strip()) for v in cells)
    if text_count < 2 or text_count / len(used) < 0.75:
        return None
    if any(re.fullmatch(r"(?:\d{4}-\d{1,2}-\d{1,2}|\d{1,2}/\d{1,2}/\d{2,4})", str(v).strip()) for v in cells if v is not None):
        return None
    data_like = sum(bool(re.fullmatch(r"[-+]?\d+(?:[.,]\d+)*|\d{4}-\d{1,2}-\d{1,2}", str(v).strip())) for v in cells if v is not None)
    if data_like >= 2 and data_like / len(used) >= .5:
        return None
    following = rows[row_index + 1:row_index + 6]
    populated = sum(any(v is not None for v in r[start:end]) for r in following)
    if populated == 0:
        return None
    headers, duplicates = _unique_headers(cells)
    warnings = []
    if duplicates:
        warnings.append("Repeated column names were numbered; confirm the intended fields.")
    if start:
        warnings.append(f"The table starts in column {start + 1}.")
    score = len(used) * 2 + populated * 2 - (row_index * .3) - (len(cells) - len(used))
    # Count actual rows when loading; this conservative estimate is only for inspection.
    return TableCandidate(sheet, row_index + 1, headers, max(0, max_row - row_index - 1), score, tuple(warnings), start + 1)


def inspect(path: str | Path) -> WorkbookInspection:
    """Inspect an xlsx or csv file, returning possible tables rather than guessing silently."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    result = WorkbookInspection(path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            rows = []
            count = 0
            for record in reader:
                count += 1
                if count <= 12:
                    rows.append(tuple(value if value.strip() else None for value in record))
        for i in range(min(5, len(rows))):
            candidate = _candidate(rows, path.stem, i, count)
            if candidate:
                result.candidates.append(candidate)
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in book:
                rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), values_only=True))
                for i in range(min(8, len(rows))):
                    candidate = _candidate(rows, sheet.title, i, sheet.max_row)
                    if candidate:
                        result.candidates.append(candidate)
        finally:
            book.close()
    else:
        raise ValueError("Supported inputs are .xlsx, .xlsm, and .csv")
    result.candidates.sort(key=lambda item: item.score, reverse=True)
    if not result.candidates:
        result.warnings.append("No table with a recognizable header and data rows was found.")
    elif len(result.candidates) > 1 and result.candidates[0].score - result.candidates[1].score < 3:
        result.warnings.append("Several tables look plausible. Choose the correct sheet and header row.")
    return result


def read_table(path: str | Path, candidate: TableCandidate) -> pd.DataFrame:
    """Read the chosen table; preserve source row numbers for later evidence links."""
    path = Path(path)
    kwargs = dict(header=None, skiprows=candidate.header_row)
    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(path, encoding="utf-8-sig", **kwargs)
    else:
        frame = pd.read_excel(path, sheet_name=candidate.sheet, engine="openpyxl", **kwargs)
    # Keep fully blank columns with headers: they are important quality findings.
    start = candidate.start_column - 1
    end = start + len(candidate.columns)
    if frame.shape[1] < end:
        frame = frame.reindex(columns=range(end))
    frame = frame.iloc[:, start:end].dropna(how="all")
    frame.columns = candidate.columns
    frame.insert(0, "_source_row", frame.index + candidate.header_row + 1)
    return frame.reset_index(drop=True)
